"""Targeted full-audit reproductions. No live Excel session is opened."""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import types
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "evals"))

# The audit only needs the harness functions; replace optional benchmark/COM
# imports so importing the module cannot launch or connect to Excel.
grader = types.ModuleType("spreadsheetbench.task_grader")
grader.compare_excels = lambda *args, **kwargs: (False, None)
package = types.ModuleType("spreadsheetbench")
sys.modules.setdefault("spreadsheetbench", package)
sys.modules["spreadsheetbench.task_grader"] = grader
client = types.ModuleType("win32com.client")
client.GetActiveObject = lambda *_: None
win32com = types.ModuleType("win32com")
win32com.client = client
sys.modules["win32com"] = win32com
sys.modules["win32com.client"] = client

import run_spreadsheetbench as runner  # noqa: E402
from workbook_diff import unauthorized_edits  # noqa: E402


def preservation_deleted_sheet(tmp: Path) -> dict:
    initial = tmp / "initial.xlsx"
    output = tmp / "output.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Answer"
    book.active["A1"] = "allowed"
    secret = book.create_sheet("SourceData")
    secret["A1"] = "must survive"
    book.save(initial)
    book.remove(book["SourceData"])
    book.save(output)
    result = unauthorized_edits(initial, output, "Answer!A1", lambda a, b: a == b)
    assert result["unauthorized_cells"] == 0
    assert result["sheets_removed"] == ["SourceData"]
    # This is the exact predicate used by summarize(): a deletion passes it.
    counted_as_damaged = bool(result["unauthorized_cells"])
    assert counted_as_damaged is False
    return {"diff": result, "countedAsDamagedBySummary": counted_as_damaged}


def open_failure_leaves_excel_flags_changed(tmp: Path) -> dict:
    source = tmp / "dataset" / "item"
    source.mkdir(parents=True)
    book = openpyxl.Workbook()
    book.save(source / "x_init.xlsx")
    shutil.copyfile(source / "x_init.xlsx", source / "x_golden.xlsx")

    class Workbooks:
        def Open(self, *_args, **_kwargs):
            raise RuntimeError("open failed")

    class App:
        DisplayAlerts = True
        Visible = False

    app = App()
    app.Workbooks = Workbooks()
    runner.ready_excel = lambda _run_dir: app
    runner.embed_taskpane = lambda *_args: None
    runner.qualified_answer_position = lambda _task: "A1"
    args = types.SimpleNamespace(model="haiku", timeout=1, compare_values=lambda a, b: a == b)
    task = {"id": "1", "instruction_type": "value", "spreadsheet_path": "item", "instruction": "x", "answer_position": "A1"}
    try:
        runner.run_task(task, tmp / "dataset", tmp / "run", args, "token", ("id", "url"), lambda *_: (False, None))
    except RuntimeError as error:
        assert str(error) == "open failed"
    else:
        raise AssertionError("expected Workbooks.Open to fail")
    assert app.DisplayAlerts is False
    assert app.Visible is True
    return {"afterOpenFailure": {"DisplayAlerts": app.DisplayAlerts, "Visible": app.Visible}}


with tempfile.TemporaryDirectory(prefix="excel-assistant-audit-") as raw:
    tmp = Path(raw)
    results = [
        {"id": "F08-deleted-sheet-not-counted-as-damage", "reproduced": True, "evidence": preservation_deleted_sheet(tmp)},
        {"id": "F09-open-failure-leaves-excel-global-state", "reproduced": True, "evidence": open_failure_leaves_excel_flags_changed(tmp)},
    ]

report = {"note": "Isolated workbook and fake COM fixtures; no live Excel session used.", "results": results}
if "--save" in sys.argv:
    (ROOT / "docs" / "full-project-audit-2026-09-22.python-repro.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
print(json.dumps(report, ensure_ascii=False, indent=2))
