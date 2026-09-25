"""Offline checks for the eval harness: run with
uv run --project evals python -m unittest discover -s evals
"""

import json
import tempfile
import unittest
from argparse import Namespace
from pathlib import Path
from unittest.mock import patch

import openpyxl

from run_spreadsheetbench import (
    classify_outcome, harness_error_record, managed_workbook, qualified_answer_position,
    run_task, select_tasks, summarize, tool_error_summary,
)
from workbook_diff import EXCEL_MAX_ROWS, authorized_ranges, unauthorized_edits
from vendor import harbor_evaluate as grader


def same_value(a, b):
    return a == b


class WorkbookDiffTest(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def save(self, name, cells, sheets=("Data",)):
        wb = openpyxl.Workbook()
        wb.active.title = sheets[0]
        for extra in sheets[1:]:
            wb.create_sheet(extra)
        for (sheet, coord), value in cells.items():
            wb[sheet][coord] = value
        path = self.dir / name
        wb.save(path)
        return path

    def test_edits_inside_authorized_range_are_ignored(self):
        before = self.save("a.xlsx", {("Data", "A1"): 1, ("Data", "B2"): None})
        after = self.save("b.xlsx", {("Data", "A1"): 1, ("Data", "B2"): "=A1*2"})
        self.assertEqual(unauthorized_edits(before, after, "B2", same_value)["unauthorized_cells"], 0)

    def test_value_and_formula_changes_outside_range_are_counted(self):
        before = self.save("a.xlsx", {("Data", "A1"): 1, ("Data", "C3"): "=A1+1"})
        after = self.save("b.xlsx", {("Data", "A1"): 5, ("Data", "C3"): 2, ("Data", "D9"): "x"})
        diff = unauthorized_edits(before, after, "B2", same_value)
        self.assertEqual(diff["unauthorized_cells"], 3)
        self.assertEqual({s["cell"] for s in diff["samples"]}, {"'Data'!A1", "'Data'!C3", "'Data'!D9"})

    def test_ranges_are_scoped_to_their_sheet(self):
        before = self.save("a.xlsx", {("Data", "A1"): 1, ("Out", "A1"): None}, sheets=("Data", "Out"))
        after = self.save("b.xlsx", {("Data", "A1"): 1, ("Out", "A1"): 9}, sheets=("Data", "Out"))
        self.assertEqual(unauthorized_edits(before, after, "'Out'!A1:B2", same_value)["unauthorized_cells"], 0)
        # Unqualified positions refer to the first sheet, as in the official grader.
        self.assertEqual(unauthorized_edits(before, after, "A1:B2", same_value)["unauthorized_cells"], 1)

    def test_added_and_removed_sheets_are_reported(self):
        before = self.save("a.xlsx", {("Data", "A1"): 1}, sheets=("Data", "Old"))
        after = self.save("b.xlsx", {("Data", "A1"): 1}, sheets=("Data", "New"))
        diff = unauthorized_edits(before, after, "A1", same_value)
        self.assertEqual((diff["sheets_removed"], diff["sheets_added"]), (["Old"], ["New"]))

    def test_removed_sheet_is_counted_as_preservation_damage(self):
        before = self.save("a.xlsx", {("Data", "A1"): 1, ("Old", "A1"): "keep"}, sheets=("Data", "Old"))
        after = self.save("b.xlsx", {("Data", "A1"): 1})
        diff = unauthorized_edits(before, after, "A1", same_value)
        result = {
            "instruction_type": "Cell", "infra_status": "ok", "agent_status": "completed",
            "tool_calls": 1, "tool_errors": 0, "agent_duration_s": 1, "passed": True,
            **diff,
            "gold_unauthorized_cells": 0, "gold_sheets_removed": [], "gold_sheets_added": [],
        }
        summary = summarize("deleted", [result])
        self.assertEqual(summary["preservation"]["tasks_with_unauthorized_edits"], 1)
        self.assertEqual(summary["preservation"]["passed_but_damaged"], 1)

    def test_authorized_ranges_parse_multi_part_positions(self):
        boxes = authorized_ranges("'My Sheet'!A1:B2,C5", "First")
        self.assertEqual(boxes, {"My Sheet": [(1, 1, 2, 2)], "First": [(5, 3, 5, 3)]})
        for position in ["Sheet1'!A1,'Sheet2'!B2", "'Sheet1!'A1,'Sheet2!'B2'", "'Sheet1'!'A1,Sheet2!B2"]:
            self.assertEqual(authorized_ranges(position, "First"), {"Sheet1": [(1, 1, 1, 1)], "Sheet2": [(2, 2, 2, 2)]})

    def test_comma_in_a_quoted_sheet_name_reaches_grading_and_preservation(self):
        # Task 130-9: the official grader split this name on its commas and crashed.
        name = "b2b, sez, de"
        position = "'b2b, sez, de'!A5:B6"
        before = self.save("a.xlsx", {(name, "A5"): 1, (name, "C9"): 1}, sheets=(name,))
        after = self.save("b.xlsx", {(name, "A5"): 2, (name, "C9"): 1}, sheets=(name,))
        self.assertEqual(authorized_ranges(position, "Other"), {name: [(5, 1, 6, 2)]})
        self.assertEqual(unauthorized_edits(before, after, position, same_value)["unauthorized_cells"], 0)
        self.assertFalse(grader.compare_workbooks(before, after, position)[0])
        self.assertTrue(grader.compare_workbooks(before, before, position)[0])

    def test_column_only_ranges_cover_every_row(self):
        # Task 283-32: "A:G" crashed the official grader; Harbor grades it to the last used row.
        position = "Sheet3'!A:B"
        self.assertEqual(authorized_ranges(position, "First"), {"Sheet3": [(1, 1, EXCEL_MAX_ROWS, 2)]})
        before = self.save("a.xlsx", {("Sheet3", "B40"): 1}, sheets=("Sheet3",))
        after = self.save("b.xlsx", {("Sheet3", "B40"): 2}, sheets=("Sheet3",))
        self.assertEqual(unauthorized_edits(before, after, position, same_value)["unauthorized_cells"], 0)
        self.assertFalse(grader.compare_workbooks(before, after, position)[0])

    def test_unqualified_parts_use_the_given_sheet_for_preservation(self):
        # Task 13-1 points the prompt at LISTS; edits there are authorized even though
        # the grader, like the benchmark's, reads an unqualified range on the first sheet.
        before = self.save("a.xlsx", {("LISTS", "A3"): 1}, sheets=("RANGES", "LISTS"))
        after = self.save("b.xlsx", {("LISTS", "A3"): 2}, sheets=("RANGES", "LISTS"))
        self.assertEqual(
            unauthorized_edits(before, after, "A3:D32", same_value, default_sheet="LISTS")["unauthorized_cells"], 0
        )
        self.assertEqual(unauthorized_edits(before, after, "A3:D32", same_value)["unauthorized_cells"], 1)


class SelectionTest(unittest.TestCase):
    tasks = [{"id": i, "instruction_type": "Cell" if i % 4 else "Sheet"} for i in range(400)]

    def test_stratified_sample_is_deterministic_and_keeps_the_mix(self):
        args = Namespace(ids=None, sample=40, seed=7, offset=0, limit=10)
        first, second = select_tasks(self.tasks, args), select_tasks(self.tasks, args)
        self.assertEqual(first, second)
        self.assertEqual(len(first), 40)
        self.assertEqual(sum(t["instruction_type"] == "Sheet" for t in first), 10)

    def test_explicit_ids_keep_the_requested_order(self):
        args = Namespace(ids=["5", "2"], sample=None, seed=0, offset=0, limit=10)
        self.assertEqual([t["id"] for t in select_tasks(self.tasks, args)], [5, 2])

    def test_answer_sheet_qualifies_every_part(self):
        task = {"answer_position": "A1:B2,D4", "answer_sheet": "Out"}
        self.assertEqual(qualified_answer_position(task), "'Out'!A1:B2,'Out'!D4")

    def test_sheet_list_is_distinct_from_a_sheet_whose_name_contains_commas(self):
        task = {"answer_position": "A1", "answer_sheet": "Output,Source"}
        self.assertEqual(qualified_answer_position(task, ["Output", "Source"]), "'Output'!A1")
        self.assertEqual(qualified_answer_position(task, ["Output,Source"]), "'Output,Source'!A1")
        self.assertEqual(qualified_answer_position({"answer_position": "A1", "answer_sheet": "Bob's"}), "'Bob''s'!A1")


class SummaryTest(unittest.TestCase):
    def test_tool_error_is_process_evidence_not_a_failure_cause(self):
        self.assertEqual(classify_outcome("ok", False, "completed"), "answer_mismatch")
        self.assertEqual(classify_outcome("ok", False, "timeout"), "agent_incomplete")
        self.assertEqual(classify_outcome("harness_error", False, None), "infrastructure")

    def test_harness_error_keeps_completed_stage_and_agent_evidence(self):
        partial = {
            "id": "120-24", "instruction_type": "Cell", "stage": "close",
            "phases_s": {"prepare": 1.2, "agent": 25.0},
            "agent_status": "completed", "tool_calls": 7, "tool_errors": 2,
        }
        result = harness_error_record(partial, RuntimeError("COM rejected close"))
        self.assertEqual(result["stage"], "close")
        self.assertEqual(result["phases_s"], partial["phases_s"])
        self.assertEqual(result["tool_calls"], 7)
        self.assertEqual(result["agent_status"], "completed")
        self.assertEqual(result["failure_class"], "infrastructure")

    def test_missing_agent_status_is_counted(self):
        record = {"instruction_type": "Cell", "infra_status": "harness_error", "agent_status": None, "passed": False}
        self.assertEqual(summarize("failed", [record])["agent_statuses"], {"None": 1})

    def test_tool_errors_are_grouped_by_root_cause(self):
        with tempfile.TemporaryDirectory() as directory:
            transcript = Path(directory) / "transcript.jsonl"
            entries = [
                {"type": "user", "message": {"content": [
                    {"type": "tool_result", "is_error": True, "content": "Input validation error: cells"},
                    {"type": "tool_result", "is_error": True, "content": "Would overwrite 2 cells"},
                ]}},
            ]
            transcript.write_text("\n".join(json.dumps(entry) for entry in entries), encoding="utf-8")
            result = tool_error_summary(transcript)
        self.assertEqual(result["count"], 2)
        self.assertEqual(result["categories"], {"argument_validation": 1, "overwrite_guard": 1})

    def test_infra_failures_stay_in_the_headline_denominator(self):
        ok = {"instruction_type": "Cell", "infra_status": "ok", "agent_status": "completed", "tool_calls": 3,
              "tool_errors": 0, "agent_duration_s": 10, "unauthorized_cells": 0}
        results = [
            {**ok, "passed": True},
            {**ok, "passed": True, "unauthorized_cells": 2},
            {**ok, "passed": False, "infra_status": "stalled"},
            {**ok, "passed": False},
            {**ok, "passed": True, "unauthorized_cells": 50, "gold_unauthorized_cells": 40},
        ]
        summary = summarize("t", results)
        self.assertEqual(summary["end_to_end_pass_rate"], 0.6)
        self.assertEqual(summary["infra_completion_rate"], 0.8)
        self.assertEqual(summary["pass_rate_given_infra_ok"], 0.75)
        self.assertEqual(summary["preservation"]["passed_but_damaged"], 1)
        self.assertEqual(summary["preservation"]["skipped_gold_edits_outside"], 1)
        self.assertIsNone(summary["agent_seconds"]["p90"])


class ManagedWorkbookTest(unittest.TestCase):
    def test_close_error_preserves_agent_work_in_partial_record(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "case"
            source.mkdir()
            wb = openpyxl.Workbook()
            wb.save(source / "case_init.xlsx")
            wb.save(source / "golden.xlsx")

            class Workbook:
                Windows = []
                def Save(self):
                    pass
                def Close(self, SaveChanges=False):
                    raise RuntimeError("close rejected")

            class App:
                DisplayAlerts = True
                Visible = False
                Workbooks = type("Workbooks", (), {"Open": lambda *_args, **_kwargs: Workbook()})()

            task = {"id": "sample", "instruction_type": "Cell", "spreadsheet_path": "case",
                    "answer_position": "A1", "instruction": "Fill A1"}
            partial = {"id": "sample", "instruction_type": "Cell", "phases_s": {}}
            agent = {"status": "completed", "tools": ["a", "b"], "monotonicMs": 3000}
            with patch("run_spreadsheetbench.embed_taskpane"), \
                 patch("run_spreadsheetbench.ready_excel", return_value=App()), \
                 patch("run_spreadsheetbench.daemon_request", return_value=agent):
                with self.assertRaisesRegex(RuntimeError, "close rejected") as caught:
                    run_task(task, root, root / "run", Namespace(model="haiku", timeout=10), "token",
                             ("manifest", "id"), lambda *_: (True, None), partial)
            result = harness_error_record(partial, caught.exception)
            self.assertEqual(result["stage"], "close")
            self.assertEqual(result["agent_status"], "completed")
            self.assertEqual(result["tool_calls"], 2)
            self.assertEqual(result["agent_duration_s"], 3.0)

    def test_open_failure_restores_excel_global_state(self):
        class Workbooks:
            @staticmethod
            def Open(*_args, **_kwargs):
                raise RuntimeError("open failed")

        class App:
            DisplayAlerts = True
            Visible = False

        app = App()
        app.Workbooks = Workbooks()
        with self.assertRaisesRegex(RuntimeError, "open failed"):
            with managed_workbook(app, "bad.xlsx"):
                pass
        self.assertTrue(app.DisplayAlerts)
        self.assertFalse(app.Visible)


if __name__ == "__main__":
    unittest.main()
